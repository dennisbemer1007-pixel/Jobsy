#!/usr/bin/env python3
"""Build sourced job-outlook JSON from the occupation cache.

No network. Reads tools/occupations/.cache/ (fetch.py) and writes:

- Jobsy.Core/Data/Occupations/outlook.json
- Jobsy.Core/Data/Occupations/occupation-skills.json
- Jobsy.Core/Data/Occupations/skills.nl.json
- tools/occupations/ilo_tasks_nl_review.csv (from ilo_tasks_nl.json)

    python tools/occupations/build_outlook.py
"""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
CACHE = ROOT / ".cache"
OUT = REPO / "Jobsy.Core" / "Data" / "Occupations"
SOURCES = ROOT / "sources.json"
TASKS_NL = ROOT / "ilo_tasks_nl.json"
REVIEW_CSV = ROOT / "ilo_tasks_nl_review.csv"

MONTHS_NL = (
    "januari",
    "februari",
    "maart",
    "april",
    "mei",
    "juni",
    "juli",
    "augustus",
    "september",
    "oktober",
    "november",
    "december",
)

SUBJECTS = {
    "ITKB toekomstige knelpunten beroepsgroep in 2030": "itkb",
    "verwachte baanopeningen tot 2030": "baanopeningen",
    "verwachte uitbreidingsvraag tot 2030": "uitbreiding",
    "verwachte vervangingsvraag tot 2030": "vervanging",
    "verwachte arbeidsmarktontwikkeling beroep 2025-2030": "richting",
    "indicator huidige arbeidsmarktsituatie beroep (2026Q1)": "huidige",
    "vergelijkbaarheidsindex": "vergelijkbaar",
}

# Coverage over the 3,039 ESCO occupations. The build fails when these move.
EXPECTED_COVERAGE = {
    "ilo": 2956,
    "roaItkbAndOpenings": 3039,
    "roaDirection": 2612,
}
EXPECTED_TASK_LINES = 1657
STATUS_OK = "goedgekeurd"
STATUS_CONCEPT = "concept, wacht op akkoord"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def dump(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")


def number(value: str | None):
    text = (value or "").strip().replace(",", ".")
    if not text:
        return None
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return None


def text_or_none(value: str | None) -> str | None:
    text = " ".join((value or "").split())
    return text or None


def repair_text(value: str) -> str:
    """Undo Mac Roman mojibake of UTF-8 apostrophes in the ILO workbook."""
    return value.replace("\u201a\u00c4\u00f4", "'").replace("\u00ac\u00a5", "'")


def brc_code(waarde: str) -> str | None:
    raw = (waarde or "").strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    if not raw.isdigit():
        return None
    if len(raw) == 5 and raw.startswith("0"):
        raw = raw[1:]
    return raw or None


def peildatum_from_versie(versie: str) -> str:
    raw = (versie or "").strip().lstrip("vV")
    if len(raw) != 8 or not raw.isdigit():
        raise SystemExit(f"unexpected ROA versie {versie!r}")
    year, month, day = int(raw[:4]), int(raw[4:6]), int(raw[6:8])
    if not 1 <= month <= 12 or not 1 <= day <= 31:
        raise SystemExit(f"unexpected ROA versie {versie!r}")
    return f"{day} {MONTHS_NL[month - 1]} {year}"


def load_catalogue() -> list[dict]:
    path = OUT / "occupations.nl.json"
    if not path.exists():
        raise SystemExit("occupations.nl.json is missing; run build_catalog.py first")
    rows = json.loads(path.read_text(encoding="utf-8"))
    if len(rows) != 3039:
        raise SystemExit(f"catalogue occupations={len(rows)}, expected 3039")
    return [
        {
            "id": row["id"],
            "nl": row.get("nl") or "",
            "isco": row.get("isco") or "",
            "brc": row.get("brc") or "",
        }
        for row in rows
    ]


def load_esco_raw() -> dict[str, dict]:
    pages = sorted((CACHE / "esco").glob("page_*.json"))
    if len(pages) != 31:
        raise SystemExit(f"expected 31 ESCO pages, found {len(pages)}. Run fetch.py.")
    by_id: dict[str, dict] = {}
    for page in pages:
        payload = json.loads(page.read_text(encoding="utf-8"))
        for raw in (payload.get("_embedded") or {}).get("results") or []:
            uri = (raw.get("uri") or "").strip()
            if not uri:
                continue
            esco_id = uri.rstrip("/").split("/")[-1].lower()
            by_id[esco_id] = raw
    if len(by_id) != 3039:
        raise SystemExit(f"ESCO occupations={len(by_id)}, expected 3039")
    return by_id


def load_brc_by_id(occupations: list[dict]) -> dict[str, str]:
    return {row["id"]: row.get("brc") or "" for row in occupations}


def load_roa() -> tuple[dict, str]:
    path = CACHE / "ais_tot_2030.csv"
    if not path.exists():
        raise SystemExit("missing ais_tot_2030.csv. Run fetch.py.")
    groups: dict[str, dict] = {}
    versions: set[str] = set()
    seen_subjects: set[str] = set()
    with path.open(encoding="cp1252", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        required = {"regionaam", "aggregatieniveau", "thema", "waarde", "onderwerp", "versie"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise SystemExit(f"ROA CSV missing columns {sorted(missing)}; header={reader.fieldnames}")
        for row in reader:
            if (row.get("regionaam") or "").strip() != "Nederland":
                continue
            level = (row.get("aggregatieniveau") or "").strip().lower()
            if "beroepsgroep" not in level:
                continue
            theme = (row.get("thema") or "").strip()
            if not theme.startswith("Risicoindicatoren"):
                continue
            subject = (row.get("onderwerp") or "").strip()
            seen_subjects.add(subject)
            key = SUBJECTS.get(subject)
            if key is None:
                continue
            code = brc_code(row.get("waarde") or "")
            if not code:
                continue
            versions.add((row.get("versie") or "").strip())
            group = groups.setdefault(code, {})
            if key == "vergelijkbaar":
                name = text_or_none(row.get("naamonderwerp"))
                if not name:
                    continue
                names = group.setdefault("vergelijkbaar", [])
                if name not in names and len(names) < 5:
                    names.append(name)
                continue
            group[key] = {
                "typering": text_or_none(row.get("typering")),
                "indicator": text_or_none(row.get("indicator")),
                "aantal": number(row.get("aantal")),
                "totaal6jrperc": number(row.get("totaal6jrperc")),
            }
    if not groups:
        sample = ", ".join(sorted(seen_subjects)[:12])
        raise SystemExit(f"ROA parser kept no beroepsgroep rows. Subjects seen: {sample}")
    versions.discard("")
    if len(versions) != 1:
        raise SystemExit(f"expected one ROA versie, found {sorted(versions)}")
    peildatum = peildatum_from_versie(next(iter(versions)))
    return {"byBrc": groups, "peildatum": peildatum}, peildatum


def load_ilo() -> dict[str, dict]:
    path = CACHE / "ilo_genai_scores.xlsx"
    if not path.exists():
        raise SystemExit("missing ilo_genai_scores.xlsx. Run fetch.py.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}
    for required in (
        "ISCO_08",
        "taskID",
        "Task_ISCO",
        "score_2025",
        "mean_score_2025",
        "SD_2025",
        "potential25",
    ):
        if required not in index:
            raise SystemExit(f"ILO sheet missing {required}: {header}")
    by_isco: dict[str, list[dict]] = defaultdict(list)
    meta: dict[str, dict] = {}
    for row in rows:
        raw_code = row[index["ISCO_08"]]
        if raw_code is None:
            continue
        if isinstance(raw_code, float):
            isco = f"{int(raw_code):04d}"
        else:
            isco = str(raw_code).strip().zfill(4)[-4:]
        if len(isco) != 4 or not isco.isdigit():
            continue
        task_id = row[index["taskID"]]
        task = row[index["Task_ISCO"]]
        score = row[index["score_2025"]]
        if task_id is None or task is None or score is None:
            continue
        by_isco[isco].append(
            {
                "taskID": str(task_id).strip(),
                "en": repair_text(" ".join(str(task).split())),
                "score": float(score),
            }
        )
        meta[isco] = {
            "potential25": " ".join(str(row[index["potential25"]] or "").split()),
            "meanScore2025": float(row[index["mean_score_2025"]]),
            "sd2025": float(row[index["SD_2025"]]),
        }
    workbook.close()
    if len(meta) != 427:
        raise SystemExit(f"ILO ISCO codes={len(meta)}, expected 427")
    result = {}
    for isco, tasks in by_isco.items():
        ordered = sorted(tasks, key=lambda item: (-item["score"], item["taskID"]))
        high = ordered[:2]
        high_ids = {item["taskID"] for item in high}
        # Keep a tied task out of both lists, so four distinct tasks stay four review lines.
        rest = [item for item in ordered if item["taskID"] not in high_ids]
        low = sorted(rest, key=lambda item: (item["score"], item["taskID"]))[:2]
        info = meta[isco]
        result[isco] = {
            "potential25": info["potential25"],
            "meanScore2025": info["meanScore2025"],
            "sd2025": info["sd2025"],
            "high": [{"taskID": item["taskID"], "en": item["en"], "score": item["score"]} for item in high],
            "low": [{"taskID": item["taskID"], "en": item["en"], "score": item["score"]} for item in low],
        }
    return result


def skill_links(raw: dict, key: str) -> list[tuple[str, str]]:
    links = (raw.get("_links") or {}).get(key) or []
    if isinstance(links, dict):
        links = [links]
    found = []
    for item in links:
        if not isinstance(item, dict):
            continue
        uri = (item.get("uri") or "").strip()
        title = item.get("title")
        if isinstance(title, dict):
            title = title.get("nl") or title.get("literal") or ""
        title = " ".join(str(title or "").split())
        if not uri or not title:
            continue
        found.append((uri, title))
    return found


def build_skills(occupations: list[dict], raw_by_id: dict[str, dict]) -> tuple[dict, list[dict]]:
    titles: dict[str, str] = {}
    per_occ: dict[str, dict[str, list[str]]] = {}
    for occ in occupations:
        raw = raw_by_id[occ["id"]]
        essential = skill_links(raw, "hasEssentialSkill")
        optional = skill_links(raw, "hasOptionalSkill")
        for uri, title in essential + optional:
            titles.setdefault(uri, title)
        per_occ[occ["id"]] = {
            "essential": [uri for uri, _ in essential],
            "optional": [uri for uri, _ in optional],
        }
    ordered = sorted(titles)
    index = {uri: pos for pos, uri in enumerate(ordered)}
    skills = [{"uri": uri, "nl": titles[uri]} for uri in ordered]
    compact = {
        esco_id: {
            "essential": [index[uri] for uri in rows["essential"] if uri in index],
            "optional": [index[uri] for uri in rows["optional"] if uri in index],
        }
        for esco_id, rows in per_occ.items()
    }
    return compact, skills


def shown_tasks(ilo: dict[str, dict], esco_iscos: set[str]) -> list[dict]:
    lines = []
    for isco in sorted(esco_iscos & set(ilo)):
        info = ilo[isco]
        seen = set()
        for bucket in ("high", "low"):
            for task in info[bucket]:
                key = task["taskID"]
                if key in seen:
                    continue
                seen.add(key)
                lines.append(
                    {
                        "isco": isco,
                        "taskID": key,
                        "en": task["en"],
                        "score": task["score"],
                    }
                )
    return lines


def example_occupation(occupations: list[dict]) -> dict[str, str]:
    by_isco: dict[str, str] = {}
    for occ in sorted(occupations, key=lambda item: (item["isco"], item["nl"].casefold(), item["id"])):
        by_isco.setdefault(occ["isco"], occ["nl"])
    return by_isco


def validate_translations(lines: list[dict]) -> list[dict]:
    if not TASKS_NL.exists():
        raise SystemExit(
            f"missing {TASKS_NL.name}. Draft the {len(lines)} task lines before building outlook."
        )
    rows = json.loads(TASKS_NL.read_text(encoding="utf-8"))
    if not isinstance(rows, list):
        raise SystemExit("ilo_tasks_nl.json must be a list")
    expected = {(line["isco"], line["taskID"]) for line in lines}
    found = {(row.get("isco"), row.get("taskID")) for row in rows}
    if found != expected:
        missing = len(expected - found)
        extra = len(found - expected)
        raise SystemExit(f"ilo task lines differ from ILO extract: missing {missing}, extra {extra}, file {len(rows)}")
    if len(rows) != EXPECTED_TASK_LINES:
        raise SystemExit(f"translation lines={len(rows)}, expected {EXPECTED_TASK_LINES}")
    by_key = {(row["isco"], row["taskID"]): row for row in rows}
    for line in lines:
        row = by_key[(line["isco"], line["taskID"])]
        status = (row.get("status") or "").strip()
        nl = (row.get("nl") or "").strip()
        if status not in {STATUS_OK, STATUS_CONCEPT, "afgewezen"}:
            raise SystemExit(f"bad translation status {status!r} for {line['isco']} {line['taskID']}")
        if status == STATUS_OK and not nl:
            raise SystemExit(f"goedgekeurd task {line['isco']} {line['taskID']} has an empty Dutch line")
        if (row.get("en") or "").strip() != line["en"]:
            raise SystemExit(f"English task drifted for {line['isco']} {line['taskID']}")
    return rows


def write_review_csv(rows: list[dict], lines: list[dict], examples: dict[str, str]) -> None:
    score = {(line["isco"], line["taskID"]): line["score"] for line in lines}
    ordered = sorted(rows, key=lambda row: (row["isco"], -score.get((row["isco"], row["taskID"]), 0), row["taskID"]))
    with REVIEW_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", lineterminator="\n")
        writer.writerow(["isco", "beroep", "taak (en)", "vertaling (nl)", "score_2025", "Akkoord (ja/nee/opmerking)"])
        for row in ordered:
            key = (row["isco"], row["taskID"])
            value = score.get(key)
            score_text = "" if value is None else f"{value:.4f}".rstrip("0").rstrip(".")
            akkoord = ""
            if row.get("status") == STATUS_OK:
                akkoord = "ja"
            elif row.get("status") == "afgewezen":
                akkoord = "nee"
            writer.writerow(
                [
                    row["isco"],
                    examples.get(row["isco"], ""),
                    row.get("en") or "",
                    row.get("nl") or "",
                    score_text,
                    akkoord,
                ]
            )


def coverage(occupations: list[dict], brc_by_id: dict[str, str], roa: dict, ilo: dict) -> dict[str, int]:
    ilo_n = 0
    itkb_open = 0
    direction = 0
    for occ in occupations:
        info = ilo.get(occ["isco"])
        if info and info.get("potential25"):
            ilo_n += 1
        group = roa["byBrc"].get(brc_by_id.get(occ["id"]) or "")
        if group:
            itkb = group.get("itkb") or {}
            openings = group.get("baanopeningen") or {}
            if itkb.get("typering") and openings.get("totaal6jrperc") is not None:
                itkb_open += 1
            direction_row = group.get("richting") or {}
            if direction_row.get("typering"):
                direction += 1
    return {"ilo": ilo_n, "roaItkbAndOpenings": itkb_open, "roaDirection": direction}


def source_entries() -> list[dict]:
    sources = json.loads(SOURCES.read_text(encoding="utf-8"))
    wanted = {
        "roa-ais-2030-2026": CACHE / "ais_tot_2030.csv",
        "roa-ais-2030-2026-toelichting": CACHE / "ais_tot_2030_toelichting.csv",
        "ilo-wp140-genai-isco08": CACHE / "ilo_genai_scores.xlsx",
        "esco-v1.2.1-nl": None,
    }
    entries = []
    for item in sources["inputs"]:
        if item["id"] not in wanted:
            continue
        entry = {
            "id": item["id"],
            "version": item["version"],
            "url": item["url"],
            "downloadedOn": sources.get("downloadedOn"),
        }
        path = wanted[item["id"]]
        if path is not None:
            entry["sha256"] = sha256_file(path)
        entries.append(entry)
    return entries


def build(out_dir: Path | None = None, write_review: bool = True) -> dict:
    out_dir = out_dir or OUT
    occupations = load_catalogue()
    brc_by_id = load_brc_by_id(occupations)
    roa, peildatum = load_roa()
    ilo = load_ilo()
    counts = coverage(occupations, brc_by_id, roa, ilo)
    print(
        "outlook coverage:"
        f" ILO potential25 {counts['ilo']} ({counts['ilo'] / 3039:.1%});"
        f" ROA ITKB + baanopeningen {counts['roaItkbAndOpenings']} ({counts['roaItkbAndOpenings'] / 3039:.1%});"
        f" ROA direction 2025-2030 {counts['roaDirection']} ({counts['roaDirection'] / 3039:.1%})"
    )
    for key, expected in EXPECTED_COVERAGE.items():
        if counts[key] != expected:
            raise SystemExit(f"coverage {key}={counts[key]}, expected {expected}")
    esco_iscos = {occ["isco"] for occ in occupations}
    lines = shown_tasks(ilo, esco_iscos)
    print(f"ILO task lines for ESCO ISCO codes: {len(lines)} ({len(esco_iscos & set(ilo))} codes)")
    if len(lines) != EXPECTED_TASK_LINES:
        raise SystemExit(f"task lines={len(lines)}, expected {EXPECTED_TASK_LINES}")
    translations = validate_translations(lines)
    if write_review:
        write_review_csv(translations, lines, example_occupation(occupations))
    skills, skill_rows = build_skills(occupations, load_esco_raw())
    # Drop task scores from the runtime file after the review CSV is written.
    ilo_public = {}
    for isco, info in sorted(ilo.items()):
        ilo_public[isco] = {
            "potential25": info["potential25"],
            "meanScore2025": info["meanScore2025"],
            "sd2025": info["sd2025"],
            "high": [{"taskID": task["taskID"], "en": task["en"]} for task in info["high"]],
            "low": [{"taskID": task["taskID"], "en": task["en"]} for task in info["low"]],
        }
    outlook = {
        "peildatum": peildatum,
        "sources": source_entries(),
        "roaByBrc": {code: roa["byBrc"][code] for code in sorted(roa["byBrc"])},
        "iloByIsco": ilo_public,
        "coverage": counts,
    }
    dump(out_dir / "outlook.json", outlook)
    dump(out_dir / "occupation-skills.json", skills)
    dump(out_dir / "skills.nl.json", skill_rows)
    print(
        f"wrote outlook ({(out_dir / 'outlook.json').stat().st_size} bytes),"
        f" skills {len(skill_rows)}, occupations {len(skills)}"
    )
    return {"coverage": counts, "taskLines": len(lines), "skills": len(skill_rows)}


def main() -> None:
    build(OUT)


if __name__ == "__main__":
    main()

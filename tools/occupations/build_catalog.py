#!/usr/bin/env python3
"""Build the committed ESCO occupation catalogue from tools/occupations/.cache/.

No network. Reads the cache written by fetch.py, plus corrections.json.
Writes Jobsy.Core/Data/Occupations/*.json and tools/occupations/last-diff.md.

    python tools/occupations/build_catalog.py
    python tools/occupations/build_catalog.py --check
"""

from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
CACHE = ROOT / ".cache"
OUT = REPO / "Jobsy.Core" / "Data" / "Occupations"
CORRECTIONS = OUT / "corrections.json"
SOURCES = ROOT / "sources.json"

LETTERS = ("R", "I", "A", "S", "E", "C")
ELEMENT_ORDER = {
    "1.B.1.a": 0,
    "1.B.1.b": 1,
    "1.B.1.c": 2,
    "1.B.1.d": 3,
    "1.B.1.e": 4,
    "1.B.1.f": 5,
}
NAME_ORDER = {
    "realistic": 0,
    "investigative": 1,
    "artistic": 2,
    "social": 3,
    "enterprising": 4,
    "conventional": 5,
}
MATCH_TIER = {
    "exactMatch": "exact",
    "narrowMatch": "narrow",
    "closeMatch": "close",
    "broadMatch": "broad",
}
TIER_RANK = {"exact": 4, "narrow": 3, "close": 2, "broad": 1}
EXPECTED_TIER = {"exact": 498, "narrow": 89, "close": 745, "broad": 1139, "none": 568}
# Research note 1712 medium / 261 low counted one occupation as low because
# 4.57 - 3.07 is 1.5000000000000004 in binary float. On the 2-decimal OI
# scale that spread is exactly 1.50, which the rule (spread <= 1.5) calls
# medium. The occupation is ouderenwerker (3b2229b4-807c-482a-89ca-2d0a542be196).
EXPECTED_CONFIDENCE = {"high": 498, "medium": 1713, "low": 260}
STATUS_OK = "goedgekeurd"
STATUS_CONCEPT = "concept, wacht op akkoord"
STATUS_REJECTED = "afgewezen"


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def as_float(value: Decimal) -> float:
    return float(value)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def git_commit() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=REPO, text=True, stderr=subprocess.DEVNULL
        ).strip()
    except (subprocess.CalledProcessError, FileNotFoundError):
        return "unknown"


def lang_text(node) -> str | None:
    if node is None:
        return None
    if isinstance(node, str):
        text = " ".join(node.split())
        return text or None
    if isinstance(node, dict):
        if "literal" in node and isinstance(node["literal"], str):
            return lang_text(node["literal"])
        for key in ("nl", "NL"):
            if key in node:
                return lang_text(node[key])
        if len(node) == 1:
            return lang_text(next(iter(node.values())))
    return None


def lang_list(node) -> list[str]:
    if node is None:
        return []
    if isinstance(node, list):
        values = []
        for item in node:
            text = lang_text(item) if not isinstance(item, str) else " ".join(item.split())
            if text:
                values.append(text)
        return values
    if isinstance(node, dict):
        for key in ("nl", "NL"):
            if key in node:
                return lang_list(node[key])
        if "literal" in node:
            text = lang_text(node)
            return [text] if text else []
    if isinstance(node, str):
        text = " ".join(node.split())
        return [text] if text else []
    return []


def load_esco() -> tuple[list[dict], str]:
    pages = sorted((CACHE / "esco").glob("page_*.json"))
    if len(pages) != 31:
        raise SystemExit(f"expected 31 ESCO pages in .cache/esco, found {len(pages)}. Run fetch.py.")
    by_uri: dict[str, dict] = {}
    for page in pages:
        payload = json.loads(page.read_text(encoding="utf-8"))
        if payload.get("total") != 3039:
            raise SystemExit(f"{page.name} total={payload.get('total')}, expected 3039")
        for raw in (payload.get("_embedded") or {}).get("results") or []:
            uri = (raw.get("uri") or "").strip()
            if not uri:
                raise SystemExit("ESCO result without uri")
            by_uri[uri] = raw
    if len(by_uri) != 3039:
        raise SystemExit(f"ESCO unique occupations={len(by_uri)}, expected 3039")
    normalised = []
    for uri in sorted(by_uri):
        raw = by_uri[uri]
        label = lang_text(raw.get("preferredLabel")) or lang_text(raw.get("title"))
        if not label:
            raise SystemExit(f"missing Dutch label for {uri}")
        description = lang_text(raw.get("description")) or ""
        code = str(raw.get("code") or "").strip()
        isco = code[:4]
        if len(isco) != 4 or not isco.isdigit():
            raise SystemExit(f"bad ISCO from code {code!r} for {uri}")
        esco_id = uri.rstrip("/").split("/")[-1].lower()
        alts = []
        seen = {label.casefold()}
        for alt in lang_list(raw.get("alternativeLabel")):
            if alt.casefold() in seen:
                continue
            seen.add(alt.casefold())
            alts.append(alt)
        normalised.append(
            {
                "id": esco_id,
                "uri": uri,
                "nl": label,
                "alt": alts,
                "desc": description.strip(),
                "isco": isco,
                "status": raw.get("status") or "",
            }
        )
    blob = json.dumps(normalised, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return normalised, sha256_bytes(blob)


def load_crosswalk() -> dict[str, list[tuple[str, str]]]:
    path = CACHE / "esco_onet_crosswalk.csv"
    if not path.exists():
        raise SystemExit("missing crosswalk cache. Run fetch.py.")
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    header_at = next((i for i, line in enumerate(lines) if line.startswith("O*NET Id")), None)
    if header_at is None:
        raise SystemExit("crosswalk header row not found")
    reader = csv.reader(lines[header_at:])
    header = next(reader)
    if not header[0].startswith("O*NET Id"):
        raise SystemExit(f"unexpected crosswalk header: {header[:3]}")
    by_esco: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for row in reader:
        if len(row) < 7:
            continue
        onet = row[0].strip()
        uri = row[3].strip()
        kind = row[6].strip()
        if kind == "exactISCO" or kind not in MATCH_TIER:
            continue
        if "/occupation/" not in uri:
            continue
        esco_id = uri.rstrip("/").split("/")[-1].lower()
        by_esco[esco_id].append((onet, MATCH_TIER[kind]))
    return by_esco


def load_oi() -> tuple[dict[str, list[Decimal]], dict[str, dict[str, str]]]:
    path = CACHE / "career_interest_types.csv"
    if not path.exists():
        raise SystemExit("missing career_interest_types.csv. Run fetch.py.")
    profiles: dict[str, list[Decimal | None]] = {}
    meta: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if (row.get("Scale ID") or "").strip() != "OI":
                continue
            code = (row.get("O*NET-SOC Code") or "").strip()
            element = (row.get("Element ID") or "").strip()
            name = (row.get("Element Name") or "").strip().lower()
            index = ELEMENT_ORDER.get(element, NAME_ORDER.get(name))
            if index is None or not code:
                continue
            slot = profiles.setdefault(code, [None] * 6)
            slot[index] = Decimal((row.get("Data Value") or "").strip())
            meta[code] = {
                "date": (row.get("Date") or "").strip(),
                "source": (row.get("Domain Source") or "").strip(),
                "title": (row.get("Title") or "").strip(),
            }
    complete: dict[str, list[Decimal]] = {}
    for code, values in profiles.items():
        if any(value is None for value in values):
            continue
        complete[code] = [value for value in values if value is not None]
    if len(complete) != 923:
        raise SystemExit(f"O*NET OI occupations={len(complete)}, expected 923")
    return complete, meta


def load_cbs() -> tuple[dict[str, dict], list[dict]]:
    path = CACHE / "brc2014.xlsx"
    if not path.exists():
        raise SystemExit("missing brc2014.xlsx. Run fetch.py.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    variables = workbook["BRC2014variabelen"]
    rows = variables.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}
    for required in ("BRC2014beroep", "BRC2014ed2025beroepsgroep", "ISCO2008niveau"):
        if required not in index:
            raise SystemExit(f"CBS sheet missing column {required}: {header}")
    by_isco: dict[str, dict] = {}
    for row in rows:
        raw_code = row[index["BRC2014beroep"]]
        if raw_code is None:
            continue
        if isinstance(raw_code, float):
            isco = f"{int(raw_code):04d}"
        else:
            isco = str(raw_code).strip().zfill(4)
        if not isco.isdigit():
            continue
        level_raw = row[index["ISCO2008niveau"]]
        level = int(level_raw) if level_raw not in (None, "") else None
        brc_raw = row[index["BRC2014ed2025beroepsgroep"]]
        brc = "" if brc_raw is None else str(brc_raw).strip()
        if brc.endswith(".0"):
            brc = brc[:-2]
        by_isco[isco] = {"brc": brc, "iscoLevel": level}

    index_sheet = workbook["BRC2014 index"]
    index_rows = index_sheet.iter_rows(values_only=True)
    index_header = [str(cell).strip() if cell is not None else "" for cell in next(index_rows)]
    index_pos = {name: pos for pos, name in enumerate(index_header)}
    title_key = "Voorbeeldberoepen_CBS"
    code_key = next(
        (
            key
            for key in index_pos
            if "BRC2014-beroep" in key.replace(" ", "") or "BRC2014beroep" in key.replace(" ", "")
        ),
        None,
    )
    if title_key not in index_pos or code_key is None:
        raise SystemExit(f"CBS index columns unexpected: {index_header}")
    titles = []
    seen = set()
    for row in index_rows:
        title = row[index_pos[title_key]]
        raw_code = row[index_pos[code_key]]
        if title is None or raw_code is None:
            continue
        if isinstance(raw_code, float):
            isco = f"{int(raw_code):04d}"
        else:
            isco = str(raw_code).strip().zfill(4)
        label = " ".join(str(title).split())
        if not label or not isco.isdigit():
            continue
        key = (label.casefold(), isco)
        if key in seen:
            continue
        seen.add(key)
        titles.append({"title": label, "isco": isco})
    workbook.close()
    return by_isco, titles


def mean_profile(codes: list[str], oi: dict[str, list[Decimal]]) -> tuple[list[Decimal], Decimal]:
    used = [oi[code] for code in codes]
    averaged = []
    spreads = []
    for index in range(6):
        column = [profile[index] for profile in used]
        averaged.append(q2(sum(column) / Decimal(len(column))))
        spreads.append(max(column) - min(column))
    return averaged, max(spreads) if spreads else Decimal(0)


def official_profile(esco_id: str, crosswalk, oi: dict[str, list[Decimal]]) -> dict:
    rows = [(code, tier) for code, tier in crosswalk.get(esco_id, []) if code in oi]
    exact = [code for code, tier in rows if tier == "exact"]
    if exact:
        used = list(dict.fromkeys(exact))
        tier = "exact"
    else:
        used = list(dict.fromkeys(code for code, _tier in rows))
        tier = "none"
        for code, row_tier in rows:
            if TIER_RANK[row_tier] > TIER_RANK.get(tier, 0):
                tier = row_tier
    if not used:
        return {"oi": None, "tier": "none", "confidence": "none", "onet": [], "spread": None}
    profile, spread = mean_profile(used, oi)
    if tier == "exact":
        confidence = "high"
    elif spread <= Decimal("1.5"):
        confidence = "medium"
    else:
        confidence = "low"
    return {
        "oi": profile,
        "tier": tier,
        "confidence": confidence,
        "onet": used,
        "spread": spread,
    }


def load_corrections() -> dict:
    if not CORRECTIONS.exists():
        raise SystemExit(f"missing {CORRECTIONS}")
    data = json.loads(CORRECTIONS.read_text(encoding="utf-8"))
    allowed = set(data.get("statusValues") or [])
    if allowed != {STATUS_CONCEPT, STATUS_OK, STATUS_REJECTED}:
        raise SystemExit(f"unexpected statusValues: {data.get('statusValues')}")
    return data


def validate_corrections(data: dict, occupations: dict[str, dict], crosswalk, oi) -> None:
    seen = set()
    for item in data["items"]:
        esco_id = item["escoId"]
        if esco_id in seen:
            raise SystemExit(f"duplicate correction {esco_id}")
        seen.add(esco_id)
        if esco_id not in occupations:
            raise SystemExit(f"correction esco id not in catalogue: {esco_id} ({item.get('nl')})")
        if item.get("status") not in {STATUS_CONCEPT, STATUS_OK, STATUS_REJECTED}:
            raise SystemExit(f"bad status on {esco_id}")
        kind = item.get("kind")
        if kind == "change":
            codes = item.get("onet") or []
            if not codes:
                raise SystemExit(f"change without onet codes: {esco_id}")
            official_codes = {code for code, _tier in crosswalk.get(esco_id, [])}
            for code in codes:
                if code not in oi:
                    raise SystemExit(f"O*NET code {code} has no OI (correction {esco_id})")
            present = all(code in official_codes for code in codes)
            if bool(item.get("inOfficialCrosswalk")) != present:
                raise SystemExit(
                    f"inOfficialCrosswalk={item.get('inOfficialCrosswalk')} but codes {codes} "
                    f"vs crosswalk {sorted(official_codes)} for {esco_id}"
                )
            if item.get("confidence") not in {"high", "medium", "low"}:
                raise SystemExit(f"change confidence missing for {esco_id}")
        elif kind == "confirm":
            if item.get("confidence") not in {"high", "medium", "low"}:
                raise SystemExit(f"confirm confidence missing for {esco_id}")
        elif kind == "search":
            hints = item.get("search") or []
            if not hints:
                raise SystemExit(f"search correction without hints: {esco_id}")
            for hint in hints:
                for key in ("alsoShow", "showFirst"):
                    for target in hint.get(key) or []:
                        if target not in occupations:
                            raise SystemExit(f"search hint {target} is not an ESCO id ({esco_id})")
        else:
            raise SystemExit(f"unknown correction kind {kind} on {esco_id}")
    for sub in data.get("levelSubstitutions") or []:
        if sub["fromEscoId"] not in occupations or sub["toEscoId"] not in occupations:
            raise SystemExit(f"level substitution ids missing: {sub}")
        if sub.get("status") not in {STATUS_CONCEPT, STATUS_OK, STATUS_REJECTED}:
            raise SystemExit(f"bad substitution status: {sub}")


def apply_approved(profile: dict, item: dict, oi: dict[str, list[Decimal]]) -> dict:
    """Return the effective profile. Official values stay on `profile`."""
    if item.get("status") != STATUS_OK:
        return profile
    kind = item["kind"]
    if kind == "search":
        return profile
    effective = dict(profile)
    if kind == "change":
        codes = list(item["onet"])
        for code in codes:
            if code not in oi:
                raise SystemExit(f"approved change uses unknown O*NET code {code}")
        averaged, _spread = mean_profile(codes, oi)
        effective["oi"] = averaged
        effective["onet"] = codes
        effective["tier"] = "lobsy"
        effective["confidence"] = item["confidence"]
    elif kind == "confirm":
        effective["tier"] = "lobsy"
        effective["confidence"] = item["confidence"]
    return effective


def counts(profiles: list[dict]) -> dict:
    tier = Counter(item["tier"] for item in profiles)
    confidence = Counter(item["confidence"] for item in profiles if item["oi"] is not None)
    return {
        "occupations": len(profiles),
        "withOi": sum(1 for item in profiles if item["oi"] is not None),
        "tier": {name: tier.get(name, 0) for name in ("exact", "narrow", "close", "broad", "none", "lobsy")},
        "confidence": {name: confidence.get(name, 0) for name in ("high", "medium", "low", "none")},
    }


def assert_official(stats: dict) -> None:
    problems = []
    if stats["occupations"] != 3039 or stats["withOi"] != 2471:
        problems.append(f"counts {stats['occupations']} / withOi {stats['withOi']}")
    for name, expected in EXPECTED_TIER.items():
        if stats["tier"].get(name) != expected:
            problems.append(f"tier {name}={stats['tier'].get(name)} expected {expected}")
    for name, expected in EXPECTED_CONFIDENCE.items():
        if stats["confidence"].get(name) != expected:
            problems.append(f"confidence {name}={stats['confidence'].get(name)} expected {expected}")
    if problems:
        print("official profile counts do not match the pinned research numbers:", file=sys.stderr)
        for problem in problems:
            print(f"  {problem}", file=sys.stderr)
        print(json.dumps(stats, indent=2), file=sys.stderr)
        raise SystemExit(1)
    print(
        f"official ok: 3039 occupations, 2471 with oi, "
        f"tier {stats['tier']}, confidence {stats['confidence']}"
    )


def dump(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")


def occupation_document(occ: dict, official: dict, effective: dict, cbs: dict | None) -> dict:
    document = {
        "id": occ["id"],
        "uri": occ["uri"],
        "nl": occ["nl"],
        "alt": occ["alt"],
        "desc": occ["desc"],
        "isco": occ["isco"],
        "iscoLevel": None if cbs is None else cbs.get("iscoLevel"),
        "brc": None if cbs is None else (cbs.get("brc") or None),
        "oi": None if effective["oi"] is None else [as_float(value) for value in effective["oi"]],
        "tier": effective["tier"],
        "confidence": effective["confidence"],
        "onet": effective["onet"],
    }
    if official["oi"] != effective["oi"]:
        document["officialOi"] = None if official["oi"] is None else [as_float(value) for value in official["oi"]]
    if official["tier"] != effective["tier"]:
        document["officialTier"] = official["tier"]
    if official["confidence"] != effective["confidence"]:
        document["officialConfidence"] = official["confidence"]
    if official["onet"] != effective["onet"]:
        document["officialOnet"] = official["onet"]
    return document


def top_letter(oi: list[Decimal] | None) -> str | None:
    if not oi:
        return None
    best = max(range(6), key=lambda index: (oi[index], -index))
    return LETTERS[best]


def write_diff(previous: list[dict] | None, current: list[dict], path: Path) -> None:
    lines = ["# Beroepencatalogus diff", ""]
    if not previous:
        lines.append("Eerste build. Er is nog geen vorige `occupations.nl.json` om mee te vergelijken.")
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return
    prev = {item["id"]: item for item in previous}
    curr = {item["id"]: item for item in current}
    added = sorted(set(curr) - set(prev))
    removed = sorted(set(prev) - set(curr))
    lines.append(f"- Toegevoegd: {len(added)}")
    lines.append(f"- Verwijderd: {len(removed)}")
    letter_changes = []
    confidence_changes = []
    for esco_id in sorted(set(prev) & set(curr)):
        before = prev[esco_id]
        after = curr[esco_id]
        if top_letter(None if before.get("oi") is None else [Decimal(str(v)) for v in before["oi"]]) != top_letter(
            None if after.get("oi") is None else [Decimal(str(v)) for v in after["oi"]]
        ):
            letter_changes.append(
                f"- {after['nl']} (`{esco_id}`): {top_letter_raw(before.get('oi'))} → {top_letter_raw(after.get('oi'))}"
            )
        if before.get("confidence") != after.get("confidence"):
            confidence_changes.append(
                f"- {after['nl']} (`{esco_id}`): {before.get('confidence')} → {after.get('confidence')}"
            )
    lines.append(f"- Hoofdletter veranderd: {len(letter_changes)}")
    lines.append(f"- Zekerheid veranderd: {len(confidence_changes)}")
    lines.append("")
    if added:
        lines.append("## Toegevoegd")
        lines.extend(f"- {curr[esco_id]['nl']} (`{esco_id}`)" for esco_id in added[:50])
        lines.append("")
    if removed:
        lines.append("## Verwijderd")
        lines.extend(f"- {prev[esco_id]['nl']} (`{esco_id}`)" for esco_id in removed[:50])
        lines.append("")
    lines.append("## Hoofdletter")
    lines.extend(letter_changes[:80] or ["- geen"])
    lines.append("")
    lines.append("## Zekerheid")
    lines.extend(confidence_changes[:80] or ["- geen"])
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def top_letter_raw(oi) -> str:
    if not oi:
        return "geen"
    return LETTERS[max(range(6), key=lambda index: (oi[index], -index))]


def build(out_dir: Path, diff_path: Path | None) -> dict:
    esco, esco_hash = load_esco()
    crosswalk = load_crosswalk()
    oi, meta = load_oi()
    cbs_by_isco, cbs_titles = load_cbs()
    corrections = load_corrections()
    by_id = {item["id"]: item for item in esco}
    validate_corrections(corrections, by_id, crosswalk, oi)
    items_by_id = {item["escoId"]: item for item in corrections["items"]}

    official_profiles = []
    documents = []
    missing_cbs = 0
    for occ in esco:
        official = official_profile(occ["id"], crosswalk, oi)
        official_profiles.append(official)
        effective = apply_approved(official, items_by_id.get(occ["id"], {}), oi)
        cbs = cbs_by_isco.get(occ["isco"])
        if cbs is None or cbs.get("iscoLevel") is None:
            missing_cbs += 1
        documents.append(occupation_document(occ, official, effective, cbs))
    if missing_cbs:
        raise SystemExit(f"{missing_cbs} occupations have no CBS skill level")

    official_stats = counts(official_profiles)
    assert_official(official_stats)
    effective_profiles = [
        {
            "oi": None if doc["oi"] is None else [Decimal(str(v)) for v in doc["oi"]],
            "tier": doc["tier"],
            "confidence": doc["confidence"],
        }
        for doc in documents
    ]
    effective_stats = counts(effective_profiles)

    source_meta = []
    sources = json.loads(SOURCES.read_text(encoding="utf-8"))
    file_for = {
        "esco-onet-crosswalk": CACHE / "esco_onet_crosswalk.csv",
        "onet-31.0-career-interest-types": CACHE / "career_interest_types.csv",
        "onet-31.0-occupation-data": CACHE / "occupation_data.csv",
        "cbs-brc-2014-ed2025": CACHE / "brc2014.xlsx",
    }
    for item in sources["inputs"]:
        entry = {
            "id": item["id"],
            "version": item["version"],
            "url": item["url"],
            "downloadedOn": sources["downloadedOn"],
        }
        if item["id"] == "esco-v1.2.1-nl":
            entry["sha256"] = esco_hash
            entry["sha256Of"] = "normalised occupation list sorted by uri"
        else:
            entry["sha256"] = sha256_file(file_for[item["id"]])
        source_meta.append(entry)

    domain_summary = Counter((row["date"], row["source"]) for row in meta.values())
    manifest = {
        "generator": "tools/occupations/build_catalog.py",
        "generatorCommit": git_commit(),
        "sources": source_meta,
        "onetDateDomainSource": [
            {"date": date, "domainSource": source, "occupations": count}
            for (date, source), count in sorted(domain_summary.items())
        ],
        "counts": {"official": official_stats, "effective": effective_stats},
        "attribution": {
            "escoEn": "This service uses the ESCO classification of the European Commission.",
            "escoNl": "Lobsy gebruikt ESCO v1.2.1 en heeft de koppeling voor een aantal beroepen aangepast.",
            "onet": (
                "This page includes information from the O*NET 31.0 Database by the U.S. Department of Labor, "
                "Employment and Training Administration (USDOL/ETA). Used under the CC BY 4.0 license. "
                "O*NET® is a trademark of USDOL/ETA. Lobsy has modified all or some of this information. "
                "USDOL/ETA has not approved, endorsed, or tested these modifications."
            ),
            "crosswalk": "ESCO–O*NET crosswalk: European Commission & U.S. Department of Labor (2022).",
            "cbs": "CBS, Beroepenclassificatie BRC 2014 editie 2025 (CC BY 4.0).",
        },
        "correctionsVersion": corrections.get("version"),
    }

    onet_doc = {
        code: {
            "oi": [as_float(q2(value)) for value in values],
            "date": meta[code]["date"],
            "source": meta[code]["source"],
            "title": meta[code]["title"],
        }
        for code, values in sorted(oi.items())
    }

    previous = None
    existing = out_dir / "occupations.nl.json"
    if existing.exists() and diff_path is not None:
        previous = json.loads(existing.read_text(encoding="utf-8"))

    dump(out_dir / "occupations.nl.json", documents)
    dump(out_dir / "cbs-title-index.json", cbs_titles)
    dump(out_dir / "manifest.json", manifest)
    dump(out_dir / "onet-oi.json", onet_doc)
    if diff_path is not None:
        write_diff(previous, documents, diff_path)
    return {"official": official_stats, "effective": effective_stats, "escoSha256": esco_hash}


def comparable(path: Path):
    data = json.loads(path.read_text(encoding="utf-8"))
    if path.name == "manifest.json" and isinstance(data, dict):
        data = dict(data)
        data.pop("generatorCommit", None)
    return data


def check() -> None:
    with tempfile.TemporaryDirectory(prefix="occ-check-") as tmp:
        out = Path(tmp)
        build(out, diff_path=None)
        names = ["occupations.nl.json", "cbs-title-index.json", "manifest.json", "onet-oi.json"]
        failed = False
        for name in names:
            committed = comparable(OUT / name)
            fresh = comparable(out / name)
            if committed != fresh:
                print(f"committed {name} differs from a rebuild", file=sys.stderr)
                failed = True
        if failed:
            raise SystemExit(1)
    print("check ok: committed catalogue matches the cache")


def main() -> None:
    if "--check" in sys.argv:
        check()
        return
    build(OUT, ROOT / "last-diff.md")


if __name__ == "__main__":
    main()
